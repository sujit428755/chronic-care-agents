"""Read any supported upload into one of two shapes:  ("table", rows)  or  ("text", plain_text).

Tables: .csv .json .xlsx   |   Documents: .txt .md .docx .pdf
The result is only a PREVIEW. Nothing is saved until the user confirms in the Data Manager.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import config


class UploadError(ValueError):
    """Raised with a plain-language message that the UI shows to the user."""


def check_upload(filename: str, size: int) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in config.ALLOWED_UPLOAD_EXTENSIONS:
        raise UploadError(f"'{ext or filename}' is not supported. Allowed: {', '.join(config.ALLOWED_UPLOAD_EXTENSIONS)}")
    if size > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise UploadError(f"File is {size / 1048576:.1f} MB; the limit is {config.MAX_UPLOAD_MB:g} MB.")
    if size == 0:
        raise UploadError("The file is empty.")
    return ext


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise UploadError("Could not read the text encoding of the file.")


def _flatten(obj, prefix="") -> dict:
    out = {}
    for k, v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, key + "_"))
        elif isinstance(v, list):
            out[key] = "; ".join(str(x) for x in v)
        else:
            out[key] = "" if v is None else v
    return out


def _json_rows(text: str) -> list[dict]:
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise UploadError(f"Not valid JSON: {exc}") from exc
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        data = lists[0] if lists else [data]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise UploadError("JSON must be a list of records (objects) or an object that contains one such list.")
    return [_flatten(r) for r in data]


def _xlsx_rows(data: bytes) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        raise UploadError("The first sheet is empty.")
    rows = []
    for r in it:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        rows.append({h: ("" if c is None else (c.strftime("%Y-%m-%d %H:%M") if hasattr(c, "strftime") and hasattr(c, "hour") else c.isoformat() if hasattr(c, "isoformat") else c))
                     for h, c in zip(header, r) if h})
    return rows


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = []
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            lines.append("")
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            lines.append("## " + txt)
        elif "list" in style:
            lines.append("- " + txt)
        else:
            lines.append(txt)
    return "\n".join(lines)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((p.extract_text() or "").strip() for p in reader.pages)
    except Exception as exc:  # noqa: BLE001 - corrupt PDFs raise many kinds of errors
        raise UploadError(f"Could not read the PDF: {exc}") from exc
    if not text.strip():
        raise UploadError("No text found in the PDF (it may be a scan). Please upload a text-based file.")
    return text


def read_upload(filename: str, data: bytes) -> tuple[str, object]:
    ext = check_upload(filename, len(data))
    if ext == ".csv":
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]
        return "table", rows
    if ext == ".json":
        return "table", _json_rows(_decode(data))
    if ext == ".xlsx":
        return "table", _xlsx_rows(data)
    if ext in (".txt", ".md"):
        return "text", _decode(data)
    if ext == ".docx":
        return "text", _docx_text(data)
    return "text", _pdf_text(data)
